"""Build the municipal access layer from official Brazilian public data.

Outputs app/data/access.json and data/access-manifest.json. Raw downloads are
kept outside version control in data/raw so every published indicator can be
reproduced and audited without shipping the source archives to the browser.
"""

from __future__ import annotations

import csv
import hashlib
import io
import json
import math
import os
import urllib.request
import zipfile
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path

import numpy as np
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
RAW = ROOT / "data" / "raw"
OUTPUT = ROOT / "app" / "data" / "access.json"
MANIFEST = ROOT / "data" / "access-manifest.json"

SOURCES = {
    "cnes": {
        "url": "https://s3.sa-east-1.amazonaws.com/ckan.saude.gov.br/CNES/cnes_estabelecimentos_csv.zip",
        "file": "cnes_estabelecimentos_csv.zip",
        "reference": "Cadastro completo disponibilizado pelo Portal de Dados Abertos do SUS",
    },
    "pfpb": {
        "url": "https://www.gov.br/saude/pt-br/composicao/sectics/farmacia-popular/credenciamento/documentacao/anexo-i-lista-de-municipios_atualizada_em_06-03-2026.xlsx/@@download/file",
        "file": "pfpb_municipios_vagas_2026.xlsx",
        "reference": "Anexo I do credenciamento do PFPB, atualizado em 06/03/2026",
    },
    "population": {
        "url": "https://apisidra.ibge.gov.br/values/t/6579/n6/all/v/9324/p/2024?formato=json",
        "file": "ibge_population_2024.json",
        "reference": "SIDRA 6579, variável 9324, população residente estimada em 2024",
    },
    "geometry": {
        "url": "https://servicodados.ibge.gov.br/api/v3/malhas/paises/BR?formato=application/vnd.geo+json&qualidade=minima&intrarregiao=municipio",
        "file": "ibge_municipios_min.geojson",
        "reference": "API de Malhas Territoriais do IBGE, qualidade mínima",
    },
}


def download(key: str) -> Path:
    RAW.mkdir(parents=True, exist_ok=True)
    item = SOURCES[key]
    target = RAW / item["file"]
    if target.exists() and target.stat().st_size:
        return target
    request = urllib.request.Request(item["url"], headers={"User-Agent": "FarmacoBrasil/0.3 academic-data-pipeline"})
    with urllib.request.urlopen(request, timeout=180) as response, target.open("wb") as output:
        while chunk := response.read(1024 * 1024):
            output.write(chunk)
    return target


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def polygon_centroid(ring: list[list[float]]) -> tuple[float, float, float]:
    area2 = cx = cy = 0.0
    for (x1, y1), (x2, y2) in zip(ring, ring[1:] + ring[:1]):
        cross = x1 * y2 - x2 * y1
        area2 += cross
        cx += (x1 + x2) * cross
        cy += (y1 + y2) * cross
    if abs(area2) < 1e-12:
        xs, ys = zip(*ring)
        return sum(xs) / len(xs), sum(ys) / len(ys), 0.0
    return cx / (3 * area2), cy / (3 * area2), abs(area2 / 2)


def geometry_centroid(geometry: dict) -> tuple[float, float]:
    polygons = [geometry["coordinates"]] if geometry["type"] == "Polygon" else geometry["coordinates"]
    weighted: list[tuple[float, float, float]] = []
    for polygon in polygons:
        if polygon and polygon[0]:
            weighted.append(polygon_centroid(polygon[0]))
    total = sum(weight for _, _, weight in weighted)
    if not weighted:
        return math.nan, math.nan
    if total == 0:
        return weighted[0][0], weighted[0][1]
    return sum(x * weight for x, _, weight in weighted) / total, sum(y * weight for _, y, weight in weighted) / total


def haversine_to_many(lon: float, lat: float, points: np.ndarray) -> np.ndarray:
    lon1, lat1 = np.radians([lon, lat])
    lon2, lat2 = np.radians(points[:, 0]), np.radians(points[:, 1])
    dlon, dlat = lon2 - lon1, lat2 - lat1
    a = np.sin(dlat / 2) ** 2 + np.cos(lat1) * np.cos(lat2) * np.sin(dlon / 2) ** 2
    return 6371.0088 * 2 * np.arctan2(np.sqrt(a), np.sqrt(1 - a))


def valid_brazil_coordinate(lon: str, lat: str) -> tuple[float, float] | None:
    try:
        x, y = float(lon), float(lat)
    except (TypeError, ValueError):
        return None
    if -75 <= x <= -32 and -35 <= y <= 6:
        return x, y
    return None


def main() -> None:
    paths = {key: download(key) for key in SOURCES}

    population_raw = json.loads(paths["population"].read_text(encoding="utf-8-sig"))
    population = {}
    for row in population_raw[1:]:
        code = str(row["D1C"])
        name, uf = row["D1N"].rsplit(" - ", 1)
        value = str(row["V"])
        population[code] = {"name": name, "uf": uf, "population": int(value) if value.isdigit() else None}

    geometry = json.loads(paths["geometry"].read_text(encoding="utf-8"))
    centroids = {
        feature["properties"]["codarea"]: geometry_centroid(feature["geometry"])
        for feature in geometry["features"]
    }

    workbook = load_workbook(paths["pfpb"], read_only=True, data_only=True)
    pfpb_rows: dict[str, int] = {}
    for row in list(workbook.active.iter_rows(values_only=True))[2:]:
        if row[2] is None:
            continue
        prefix = str(int(row[2])).zfill(6)
        pfpb_rows[prefix] = int(row[5] or 0)

    cnes_count: defaultdict[str, int] = defaultdict(int)
    cnes_sus: defaultdict[str, int] = defaultdict(int)
    cnes_points: list[tuple[float, float, str]] = []
    with zipfile.ZipFile(paths["cnes"]) as archive:
        filename = archive.namelist()[0]
        stream = io.TextIOWrapper(archive.open(filename), encoding="cp1252", newline="")
        for row in csv.DictReader(stream, delimiter=";"):
            if (row.get("TP_UNIDADE") or "").strip() != "43":
                continue
            if (row.get("CO_MOTIVO_DESAB") or "").strip():
                continue
            prefix = (row.get("CO_IBGE") or "").strip().zfill(6)
            if len(prefix) != 6:
                continue
            cnes_count[prefix] += 1
            if (row.get("CO_AMBULATORIAL_SUS") or "").strip().upper() == "SIM":
                cnes_sus[prefix] += 1
            point = valid_brazil_coordinate(row.get("NU_LONGITUDE", ""), row.get("NU_LATITUDE", ""))
            if point:
                cnes_points.append((point[0], point[1], prefix))

    municipality_name_by_prefix = {code[:6]: values["name"] for code, values in population.items()}
    covered_codes = []
    for code in population:
        filled = pfpb_rows.get(code[:6])
        if filled is None or filled > 0:
            covered_codes.append(code)

    covered_codes_with_geometry = [code for code in covered_codes if code in centroids]
    covered_coordinates = np.array([centroids[code] for code in covered_codes_with_geometry], dtype=float)
    pharmacy_coordinates = np.array([(lon, lat) for lon, lat, _ in cnes_points], dtype=float)

    municipalities = []
    for code, values in population.items():
        prefix = code[:6]
        lon, lat = centroids.get(code, (math.nan, math.nan))
        count = cnes_count[prefix]
        rate = count / values["population"] * 10000 if values["population"] else None
        filled = pfpb_rows.get(prefix)
        pfpb_covered = filled is None or filled > 0

        nearest_cnes_km = 0.0 if count else None
        nearest_cnes_name = values["name"]
        if count == 0 and len(pharmacy_coordinates) and math.isfinite(lon):
            distances = haversine_to_many(lon, lat, pharmacy_coordinates)
            index = int(np.argmin(distances))
            nearest_cnes_km = float(distances[index])
            nearest_cnes_name = municipality_name_by_prefix.get(cnes_points[index][2], "não identificado")

        nearest_pfpb_km = 0.0 if pfpb_covered else None
        nearest_pfpb_name = values["name"] if pfpb_covered else None
        if not pfpb_covered and len(covered_coordinates) and math.isfinite(lon):
            distances = haversine_to_many(lon, lat, covered_coordinates)
            index = int(np.argmin(distances))
            nearest_pfpb_km = float(distances[index])
            nearest_pfpb_name = population[covered_codes_with_geometry[index]]["name"]

        municipalities.append({
            "id": code,
            "name": values["name"],
            "uf": values["uf"],
            "population": values["population"],
            "cnesPharmacies": count,
            "cnesRate": round(rate, 2) if rate is not None else None,
            "pfpbCovered": pfpb_covered,
            "nearestCnesKm": round(nearest_cnes_km, 1) if nearest_cnes_km is not None else None,
            "nearestCnesMunicipality": nearest_cnes_name if count == 0 else None,
            "nearestPfpbKm": round(nearest_pfpb_km, 1) if nearest_pfpb_km is not None else None,
            "nearestPfpbMunicipality": nearest_pfpb_name if not pfpb_covered else None,
        })

    uf_summary = []
    for uf in sorted({item["uf"] for item in municipalities}):
        group = [item for item in municipalities if item["uf"] == uf]
        pop = sum(item["population"] or 0 for item in group)
        pharmacies = sum(item["cnesPharmacies"] for item in group)
        covered = sum(1 for item in group if item["pfpbCovered"])
        uf_summary.append({
            "uf": uf,
            "municipalities": len(group),
            "population": pop,
            "cnesPharmacies": pharmacies,
            "cnesRate": round(pharmacies / pop * 10000, 2),
            "pfpbCovered": covered,
            "pfpbCoveragePct": round(covered / len(group) * 100, 1),
        })

    total_population = sum(item["population"] or 0 for item in municipalities)
    total_pharmacies = sum(item["cnesPharmacies"] for item in municipalities)
    total_covered = sum(1 for item in municipalities if item["pfpbCovered"])
    municipality_fields = ["id", "name", "uf", "population", "cnesPharmacies", "cnesRate", "pfpbCovered", "nearestCnesKm", "nearestCnesMunicipality", "nearestPfpbKm", "nearestPfpbMunicipality"]
    payload = {
        "meta": {
            "generatedAt": datetime.now(timezone.utc).isoformat(),
            "populationYear": 2024,
            "pfpbReference": "06/03/2026",
            "cnesRetrieved": "24/08/2026",
            "municipalities": len(municipalities),
            "population": total_population,
            "cnesPharmacies": total_pharmacies,
            "cnesWithCoordinates": len(cnes_points),
            "cnesRate": round(total_pharmacies / total_population * 10000, 2),
            "pfpbCoveredMunicipalities": total_covered,
            "pfpbCoveragePct": round(total_covered / len(municipalities) * 100, 1),
            "pfpbNoCoverageObserved": len(municipalities) - total_covered,
            "method": "CNES ativo tipo 43; cobertura PFPB inferida pelo Anexo I do edital; distâncias geodésicas em linha reta entre ponto/centróide ou centróides municipais.",
        },
        "municipalityFields": municipality_fields,
        "ufs": uf_summary,
        "municipalities": [
            [item[field] for field in municipality_fields]
            for item in sorted(municipalities, key=lambda item: (item["uf"], item["name"]))
        ],
    }

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT.write_text(json.dumps(payload, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")

    manifest = {
        "generated_at": payload["meta"]["generatedAt"],
        "output": str(OUTPUT.relative_to(ROOT)).replace(os.sep, "/"),
        "rules": [
            "CNES: apenas registros ativos com TP_UNIDADE=43 (FARMACIA).",
            "PFPB: município fora do Anexo I ou com vaga preenchida é classificado como cobertura observada; a regra é uma inferência documentada.",
            "Taxas: número de farmácias CNES por 10 mil habitantes, usando população estimada de 2024.",
            "Distâncias: linha reta; zero significa presença observada no próprio município; não equivale a tempo rodoviário.",
            "Nenhum indicador confirma estoque ou dispensação efetiva.",
        ],
        "sources": {
            key: {**value, "sha256": sha256(paths[key]), "bytes": paths[key].stat().st_size}
            for key, value in SOURCES.items()
        },
    }
    MANIFEST.parent.mkdir(parents=True, exist_ok=True)
    MANIFEST.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(payload["meta"], ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
